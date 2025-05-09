import os
import json
import customtkinter as ctk
from tkinter import filedialog, messagebox, Toplevel
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PIL import Image

# Definindo cores personalizadas em tons de lilás pastel
CORES = {
    "primaria": "#D8BFD8",  # Lilás pastel claro
    "secundaria": "#E6E6FA",  # Lavanda muito claro
    "terciaria": "#C9A0DC",  # Lilás médio
    "texto": "#4B0082",  # Índigo (para texto)
    "botao": "#9370DB",  # Lilás médio para botões
    "botao_hover": "#8A2BE2",  # Violeta para hover
    "estrela": "#9932CC",  # Orquídea escura para estrelas
    "fundo": "#F8F4FF"  # Fundo muito claro com tom lilás
}

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")  # Mantemos o tema base, mas vamos sobrescrever as cores


class Livro:
    def __init__(self, titulo, autor, status_leitura, nota, tipo, imagem_path=None):
        self.titulo = titulo
        self.autor = autor
        self.status_leitura = status_leitura  # "Lido", "Não lido" ou "Lendo"
        self.nota = nota
        self.tipo = tipo
        self.imagem_path = imagem_path

    def to_dict(self):
        return {
            "titulo": self.titulo,
            "autor": self.autor,
            "status_leitura": self.status_leitura,
            "nota": self.nota,
            "tipo": self.tipo,
            "imagem_path": self.imagem_path
        }


class BibliotecaApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("The Bookshelf")
        self.geometry("800x650")
        self.configure(fg_color=CORES["fundo"])  # Cor de fundo personalizada
        self.livros = []
        self.imagem_path = None
        self.imagem_refs = {}  # Para manter referência das imagens exibidas na lista
        
        # Configurar ícone da aplicação se disponível
        try:
            self.iconbitmap("bookshelf_icon.ico")
        except:
            pass  # Ignora se o ícone não existir
            
        self.setup_ui()
        self.load_livros()

    def setup_ui(self):
        # Cabeçalho com título e subtítulo
        header_frame = ctk.CTkFrame(self, fg_color=CORES["primaria"], corner_radius=0)
        header_frame.pack(fill="x", pady=(0, 20))
        
        # Título da aplicação
        titulo_app = ctk.CTkLabel(
            header_frame, 
            text="The Bookshelf", 
            font=ctk.CTkFont(size=28, weight="bold"),
            text_color=CORES["texto"]
        )
        titulo_app.pack(pady=(20, 5))
        
        subtitulo_app = ctk.CTkLabel(
            header_frame, 
            text="Sua biblioteca pessoal", 
            font=ctk.CTkFont(size=16),
            text_color=CORES["texto"]
        )
        subtitulo_app.pack(pady=(0, 20))

        # Área principal com barra lateral e conteúdo
        main_frame = ctk.CTkFrame(self, fg_color="transparent")
        main_frame.pack(padx=15, pady=10, fill="both", expand=True)
        
        # Barra lateral com botões de ação
        sidebar_frame = ctk.CTkFrame(main_frame, fg_color=CORES["secundaria"], width=180)
        sidebar_frame.pack(side="left", fill="y", padx=(0, 15), pady=5)
        sidebar_frame.pack_propagate(False)  # Impede que o frame encolha
        
        # Título da barra lateral
        sidebar_title = ctk.CTkLabel(
            sidebar_frame, 
            text="Ações", 
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=CORES["texto"]
        )
        sidebar_title.pack(pady=(20, 15))
        
        # Botão de adicionar livro
        btn_adicionar = ctk.CTkButton(
            sidebar_frame, 
            text="Adicionar Livro", 
            command=self.abrir_janela_adicionar,
            height=40,
            corner_radius=8,
            fg_color=CORES["botao"],
            hover_color=CORES["botao_hover"],
            text_color="white",
            font=ctk.CTkFont(size=14, weight="bold")
        )
        btn_adicionar.pack(padx=15, pady=10, fill="x")
        
        # Separador
        separator = ctk.CTkFrame(sidebar_frame, height=1, fg_color=CORES["terciaria"])
        separator.pack(fill="x", padx=15, pady=15)
        
        # Título da seção de exportação
        export_title = ctk.CTkLabel(
            sidebar_frame, 
            text="Exportar", 
            font=ctk.CTkFont(size=16, weight="bold"),
            text_color=CORES["texto"]
        )
        export_title.pack(pady=(5, 10))
        
        # Botões de exportação
        ctk.CTkButton(
            sidebar_frame, 
            text="Exportar PDF", 
            command=self.exportar_pdf,
            fg_color=CORES["botao"],
            hover_color=CORES["botao_hover"],
            text_color="white"
        ).pack(padx=15, pady=5, fill="x")
        
        ctk.CTkButton(
            sidebar_frame, 
            text="Exportar Excel", 
            command=self.exportar_excel,
            fg_color=CORES["botao"],
            hover_color=CORES["botao_hover"],
            text_color="white"
        ).pack(padx=15, pady=5, fill="x")
        
        ctk.CTkButton(
            sidebar_frame, 
            text="Exportar Word", 
            command=self.exportar_word,
            fg_color=CORES["botao"],
            hover_color=CORES["botao_hover"],
            text_color="white"
        ).pack(padx=15, pady=5, fill="x")

        # Frame para mostrar lista de livros com scroll
        content_frame = ctk.CTkFrame(main_frame, fg_color=CORES["secundaria"])
        content_frame.pack(side="right", fill="both", expand=True, pady=5)
        
        # Título da lista de livros
        list_title = ctk.CTkLabel(
            content_frame, 
            text="Meus Livros", 
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=CORES["texto"]
        )
        list_title.pack(pady=(15, 10))
        
        # Frame para a lista com scroll
        self.lista_livros_frame = ctk.CTkFrame(content_frame, fg_color="transparent")
        self.lista_livros_frame.pack(padx=10, pady=10, fill="both", expand=True)

        self.canvas = ctk.CTkCanvas(self.lista_livros_frame, borderwidth=0, highlightthickness=0, bg=CORES["secundaria"])
        self.scrollbar = ctk.CTkScrollbar(self.lista_livros_frame, orientation="vertical", command=self.canvas.yview)
        self.scrollable_frame = ctk.CTkFrame(self.canvas, fg_color="transparent")

        self.scrollable_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )

        self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

    def abrir_janela_adicionar(self):
        # Janela para adicionar novo livro
        window = Toplevel(self)
        window.title("Adicionar Novo Livro")
        window.geometry("700x500")
        window.grab_set()
        
        # Estilizar a janela
        frame_principal = ctk.CTkFrame(window, fg_color=CORES["fundo"])
        frame_principal.pack(padx=0, pady=0, fill="both", expand=True)
        
        # Cabeçalho com título
        header_frame = ctk.CTkFrame(frame_principal, fg_color=CORES["primaria"], corner_radius=0, height=80)
        header_frame.pack(fill="x")
        header_frame.pack_propagate(False)  # Mantém a altura fixa
        
        # Título da janela
        titulo_janela = ctk.CTkLabel(
            header_frame, 
            text="Adicionar Novo Livro", 
            font=ctk.CTkFont(size=22, weight="bold"),
            text_color=CORES["texto"]
        )
        titulo_janela.pack(pady=25)
        
        # Layout em duas colunas
        frame_conteudo = ctk.CTkFrame(frame_principal)
        frame_conteudo.pack(padx=10, pady=5, fill="both", expand=True)
        
        # Coluna esquerda - Informações básicas
        frame_esquerda = ctk.CTkFrame(frame_conteudo, fg_color="transparent")
        frame_esquerda.pack(side="left", fill="both", expand=True, padx=10)
        
        # Coluna direita - Avaliação e imagem
        frame_direita = ctk.CTkFrame(frame_conteudo, fg_color="transparent")
        frame_direita.pack(side="right", fill="both", expand=True, padx=10)
        
        # === COLUNA ESQUERDA ===
        
        # Título - com label e campo na mesma linha
        frame_titulo = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_titulo.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_titulo, text="Título:", width=80, anchor="w").pack(side="left")
        titulo_input = ctk.CTkEntry(frame_titulo, placeholder_text="Título do livro")
        titulo_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Autor
        frame_autor = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_autor.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_autor, text="Autor:", width=80, anchor="w").pack(side="left")
        autor_input = ctk.CTkEntry(frame_autor, placeholder_text="Nome do autor")
        autor_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Tipo
        frame_tipo = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_tipo.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_tipo, text="Tipo:", width=80, anchor="w").pack(side="left")
        tipo_input = ctk.CTkComboBox(frame_tipo, values=["Físico", "Digital", "E-book", "Audiobook"])
        tipo_input.set("Físico")
        tipo_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Status de leitura
        frame_status = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_status.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_status, text="Status:", width=80, anchor="w").pack(side="left")
        status_leitura_input = ctk.CTkComboBox(frame_status, values=["Lido", "Não lido", "Lendo"])
        status_leitura_input.set("Não lido")
        status_leitura_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # === COLUNA DIREITA ===
        
        # Avaliação com estrelas
        frame_avaliacao = ctk.CTkFrame(frame_direita, fg_color="transparent")
        frame_avaliacao.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_avaliacao, text="Avaliação:", anchor="w").pack(fill="x")
        
        frame_estrelas = ctk.CTkFrame(frame_avaliacao, fg_color="transparent")
        frame_estrelas.pack(fill="x", pady=(5, 0))
        
        nota_var = ctk.IntVar(value=0)
        
        def selecionar_nota(valor):
            nota_var.set(valor)
            for i in range(1, 6):
                if i <= valor:
                    estrelas_btn[i-1].configure(text="★", font=ctk.CTkFont(size=20))
                else:
                    estrelas_btn[i-1].configure(text="☆", font=ctk.CTkFont(size=20))
        
        estrelas_btn = []
        for i in range(1, 6):
            btn = ctk.CTkButton(
                frame_estrelas, 
                text="☆", 
                width=25, 
                height=25, 
                corner_radius=12,
                font=ctk.CTkFont(size=20),
                command=lambda v=i: selecionar_nota(v),
                fg_color="transparent", 
                hover_color="#e0e0e0",
                text_color="#FFD700"
            )
            btn.pack(side="left", padx=2)
            estrelas_btn.append(btn)
        
        # Seleção de imagem
        frame_imagem_titulo = ctk.CTkFrame(frame_direita, fg_color="transparent")
        frame_imagem_titulo.pack(fill="x", pady=(15, 5))
        ctk.CTkLabel(frame_imagem_titulo, text="Imagem da capa:", anchor="w").pack(fill="x")
        
        frame_imagem = ctk.CTkFrame(frame_direita)
        frame_imagem.pack(fill="x")
        
        imagem_path_var = ctk.StringVar()
        label_imagem = ctk.CTkLabel(frame_imagem, text="Nenhuma imagem selecionada", height=60)
        label_imagem.pack(side="left", padx=10, fill="x", expand=True)
        
        def selecionar_nova_imagem():
            path = filedialog.askopenfilename(filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.gif")])
            if path:
                imagem_path_var.set(path)
                label_imagem.configure(text=os.path.basename(path))
        
        btn_imagem = ctk.CTkButton(frame_imagem, text="Selecionar", command=selecionar_nova_imagem)
        btn_imagem.pack(side="right", padx=10)
        
        # Botões de ação
        frame_acoes = ctk.CTkFrame(frame_principal, fg_color="transparent")
        frame_acoes.pack(pady=10, fill="x")
        
        def salvar_livro():
            titulo = titulo_input.get()
            autor = autor_input.get()
            
            if not titulo or not autor:
                messagebox.showwarning("Erro", "Título e autor são obrigatórios.")
                return
                
            livro = Livro(
                titulo=titulo,
                autor=autor,
                status_leitura=status_leitura_input.get(),
                nota=nota_var.get(),
                tipo=tipo_input.get(),
                imagem_path=imagem_path_var.get() if imagem_path_var.get() else None
            )
            
            self.livros.append(livro)
            self.salvar_livros()
            self.atualizar_lista()
            window.destroy()
        
        btn_cancelar = ctk.CTkButton(
            frame_acoes, 
            text="Cancelar", 
            fg_color="#f0f0f0", 
            text_color="black",
            hover_color="#e0e0e0",
            command=window.destroy
        )
        btn_cancelar.pack(side="left", padx=20, expand=True, fill="x")
        
        btn_salvar = ctk.CTkButton(
            frame_acoes, 
            text="Salvar", 
            command=salvar_livro
        )
        btn_salvar.pack(side="left", padx=20, expand=True, fill="x")

    def abrir_edicao(self, index):
        livro = self.livros[index]

        # Janela de edição
        window = Toplevel(self)
        window.title(f"Editar livro: {livro.titulo}")
        window.geometry("700x450")
        window.grab_set()
        
        # Estilizar a janela
        frame_principal = ctk.CTkFrame(window)
        frame_principal.pack(padx=20, pady=20, fill="both", expand=True)
        
        # Título da janela
        titulo_janela = ctk.CTkLabel(frame_principal, text=f"Editar Livro", font=ctk.CTkFont(size=20, weight="bold"))
        titulo_janela.pack(pady=(10, 20))
        
        # Layout em duas colunas
        frame_conteudo = ctk.CTkFrame(frame_principal)
        frame_conteudo.pack(padx=10, pady=5, fill="both", expand=True)
        
        # Coluna esquerda - Informações básicas
        frame_esquerda = ctk.CTkFrame(frame_conteudo, fg_color="transparent")
        frame_esquerda.pack(side="left", fill="both", expand=True, padx=10)
        
        # Coluna direita - Avaliação e imagem
        frame_direita = ctk.CTkFrame(frame_conteudo, fg_color="transparent")
        frame_direita.pack(side="right", fill="both", expand=True, padx=10)
        
        # === COLUNA ESQUERDA ===
        
        # Título - com label e campo na mesma linha
        frame_titulo = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_titulo.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_titulo, text="Título:", width=80, anchor="w").pack(side="left")
        titulo_input = ctk.CTkEntry(frame_titulo, placeholder_text="Título do livro")
        titulo_input.insert(0, livro.titulo)
        titulo_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Autor
        frame_autor = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_autor.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_autor, text="Autor:", width=80, anchor="w").pack(side="left")
        autor_input = ctk.CTkEntry(frame_autor, placeholder_text="Nome do autor")
        autor_input.insert(0, livro.autor)
        autor_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Tipo
        frame_tipo = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_tipo.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_tipo, text="Tipo:", width=80, anchor="w").pack(side="left")
        tipo_input = ctk.CTkComboBox(frame_tipo, values=["Físico", "Digital", "E-book", "Audiobook"])
        tipo_input.set(livro.tipo)
        tipo_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # Status de leitura
        frame_status = ctk.CTkFrame(frame_esquerda, fg_color="transparent")
        frame_status.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_status, text="Status:", width=80, anchor="w").pack(side="left")
        status_leitura_input = ctk.CTkComboBox(frame_status, values=["Lido", "Não lido", "Lendo"])
        status_leitura_input.set(livro.status_leitura if hasattr(livro, 'status_leitura') else "Lido" if getattr(livro, 'lido', False) else "Não lido")
        status_leitura_input.pack(side="left", fill="x", expand=True, padx=(5, 0))
        
        # === COLUNA DIREITA ===
        
        # Avaliação com estrelas
        frame_avaliacao = ctk.CTkFrame(frame_direita, fg_color="transparent")
        frame_avaliacao.pack(fill="x", pady=5)
        ctk.CTkLabel(frame_avaliacao, text="Avaliação:", anchor="w").pack(fill="x")
        
        frame_estrelas = ctk.CTkFrame(frame_avaliacao, fg_color="transparent")
        frame_estrelas.pack(fill="x", pady=(5, 0))
        
        nota_var = ctk.IntVar(value=livro.nota)
        
        def selecionar_nota(valor):
            nota_var.set(valor)
            for i in range(1, 6):
                if i <= valor:
                    estrelas_btn[i-1].configure(text="★", font=ctk.CTkFont(size=20))
                else:
                    estrelas_btn[i-1].configure(text="☆", font=ctk.CTkFont(size=20))
        
        estrelas_btn = []
        for i in range(1, 6):
            btn = ctk.CTkButton(
                frame_estrelas, 
                text="☆" if i > livro.nota else "★", 
                width=25, 
                height=25, 
                corner_radius=12,
                font=ctk.CTkFont(size=20),
                command=lambda v=i: selecionar_nota(v),
                fg_color="transparent", 
                hover_color="#e0e0e0",
                text_color="#FFD700"
            )
            btn.pack(side="left", padx=2)
            estrelas_btn.append(btn)
        
        # Seleção de imagem
        frame_imagem_titulo = ctk.CTkFrame(frame_direita, fg_color="transparent")
        frame_imagem_titulo.pack(fill="x", pady=(15, 5))
        ctk.CTkLabel(frame_imagem_titulo, text="Imagem da capa:", anchor="w").pack(fill="x")
        
        frame_imagem = ctk.CTkFrame(frame_direita)
        frame_imagem.pack(fill="x")
        
        imagem_path_var = ctk.StringVar(value=livro.imagem_path if livro.imagem_path else "")
        label_imagem = ctk.CTkLabel(
            frame_imagem, 
            text=os.path.basename(imagem_path_var.get()) if imagem_path_var.get() else "Nenhuma imagem selecionada",
            height=60
        )
        label_imagem.pack(side="left", padx=10, fill="x", expand=True)
        
        def selecionar_nova_imagem():
            path = filedialog.askopenfilename(filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.gif")])
            if path:
                imagem_path_var.set(path)
                label_imagem.configure(text=os.path.basename(path))
        
        btn_imagem = ctk.CTkButton(frame_imagem, text="Selecionar", command=selecionar_nova_imagem)
        btn_imagem.pack(side="right", padx=10)
        
        # Botões de ação
        frame_acoes = ctk.CTkFrame(frame_principal, fg_color="transparent")
        frame_acoes.pack(pady=10, fill="x")
        
        def salvar_modificacoes():
            livro.titulo = titulo_input.get()
            livro.autor = autor_input.get()
            livro.status_leitura = status_leitura_input.get()
            # Compatibilidade com versões anteriores
            if hasattr(livro, 'lido'):
                del livro.lido
            livro.nota = nota_var.get()
            livro.tipo = tipo_input.get()
            livro.imagem_path = imagem_path_var.get() if imagem_path_var.get() else None
            self.salvar_livros()
            self.atualizar_lista()
            window.destroy()
        
        def excluir_livro():
            resposta = messagebox.askyesno("Confirmar Exclusão", f"Deseja realmente excluir o livro '{livro.titulo}'?")
            if resposta:
                del self.livros[index]
                self.salvar_livros()
                self.atualizar_lista()
                window.destroy()
        
        frame_botoes = ctk.CTkFrame(frame_acoes, fg_color="transparent")
        frame_botoes.pack(side="left", expand=True, fill="x")
        
        btn_cancelar = ctk.CTkButton(
            frame_botoes, 
            text="Cancelar", 
            fg_color="#f0f0f0", 
            text_color="black",
            hover_color="#e0e0e0",
            command=window.destroy
        )
        btn_cancelar.pack(side="left", padx=5, expand=True, fill="x")
        
        btn_salvar = ctk.CTkButton(
            frame_botoes, 
            text="Salvar", 
            command=salvar_modificacoes
        )
        btn_salvar.pack(side="left", padx=5, expand=True, fill="x")
        
        # Botão de excluir separado
        btn_excluir = ctk.CTkButton(
            frame_acoes, 
            text="Excluir Livro", 
            fg_color="#FF5555", 
            hover_color="#FF0000",
            command=excluir_livro
        )
        btn_excluir.pack(side="right", padx=20)

    def atualizar_lista(self):
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        self.imagem_refs.clear()

        if not self.livros:
            # Mensagem quando não há livros
            msg_frame = ctk.CTkFrame(self.scrollable_frame, fg_color="transparent")
            msg_frame.pack(fill="both", expand=True, padx=20, pady=50)
            
            msg = ctk.CTkLabel(
                msg_frame, 
                text="Sua biblioteca está vazia.\nClique em 'Adicionar Livro' para começar!",
                font=ctk.CTkFont(size=16),
                justify="center"
            )
            msg.pack(pady=50)
            return

        for idx, livro in enumerate(self.livros):
            frame_item = ctk.CTkFrame(self.scrollable_frame, corner_radius=10)
            frame_item.pack(fill="x", padx=10, pady=8)

            # Imagem do livro (thumbnail)
            if livro.imagem_path and os.path.exists(livro.imagem_path):
                try:
                    img = Image.open(livro.imagem_path).convert("RGBA")
                    img.thumbnail((80, 120), Image.LANCZOS)
                    photo = ctk.CTkImage(light_image=img, dark_image=img, size=img.size)
                    self.imagem_refs[idx] = photo
                    label_img = ctk.CTkLabel(frame_item, image=photo, text=None)
                    label_img.pack(side="left", padx=10, pady=10)
                except Exception as e:
                    print(f"Erro ao carregar imagem do livro: {e}")
                    spacer = ctk.CTkLabel(frame_item, width=80)
                    spacer.pack(side="left", padx=10, pady=10)
            else:
                spacer = ctk.CTkLabel(frame_item, width=80)
                spacer.pack(side="left", padx=10, pady=10)

            # Texto do livro (título, autor, status)
            texto_frame = ctk.CTkFrame(frame_item, fg_color="transparent")
            texto_frame.pack(side="left", fill="both", expand=True, padx=5, pady=10)

            label_titulo = ctk.CTkLabel(texto_frame, text=livro.titulo, font=ctk.CTkFont(size=16, weight="bold"), anchor="w")
            label_titulo.pack(fill="x")

            label_autor = ctk.CTkLabel(texto_frame, text=f"Autor: {livro.autor}", font=ctk.CTkFont(size=12), anchor="w")
            label_autor.pack(fill="x", pady=(4, 0))
            
            # Exibir estrelas para a nota
            estrelas_frame = ctk.CTkFrame(texto_frame, fg_color="transparent")
            estrelas_frame.pack(fill="x", pady=(4, 0))
            
            nota_texto = "".join(["★" if i < livro.nota else "☆" for i in range(5)])
            label_estrelas = ctk.CTkLabel(
                estrelas_frame, 
                text=nota_texto, 
                font=ctk.CTkFont(size=14),
                text_color="#FFD700",
                anchor="w"
            )
            label_estrelas.pack(side="left")

            # Determinar o status de leitura
            status_texto = livro.status_leitura if hasattr(livro, 'status_leitura') else "Lido" if getattr(livro, 'lido', False) else "Não lido"
            
            # Ícones para status de leitura
            status_icon = {
                "Lido": "✔️",
                "Não lido": "❌",
                "Lendo": "📖"
            }.get(status_texto, "❓")
            
            label_status = ctk.CTkLabel(texto_frame, text=f"Status: {status_icon} {status_texto} | Tipo: {livro.tipo}", font=ctk.CTkFont(size=12), anchor="w")
            label_status.pack(fill="x", pady=(4, 0))

            # Botão Editar no canto direito
            btn_editar = ctk.CTkButton(frame_item, text="Editar", width=60, command=lambda i=idx: self.abrir_edicao(i))
            btn_editar.pack(side="right", padx=10, pady=10)

    def salvar_livros(self):
        with open("biblioteca.json", "w", encoding="utf-8") as f:
            json.dump([livro.to_dict() for livro in self.livros], f, ensure_ascii=False, indent=2)

    def load_livros(self):
        if os.path.exists("biblioteca.json"):
            with open("biblioteca.json", "r", encoding="utf-8") as f:
                dados = json.load(f)
                for livro_data in dados:
                    # Compatibilidade com versões anteriores
                    if 'status_leitura' not in livro_data and 'lido' in livro_data:
                        livro_data['status_leitura'] = "Lido" if livro_data['lido'] else "Não lido"
                    if 'tipo' not in livro_data:
                        livro_data['tipo'] = 'Físico'
                self.livros = []
                for d in dados:
                    # Remover campo 'lido' se existir para evitar duplicação
                    if 'lido' in d and 'status_leitura' in d:
                        d.pop('lido')
                    self.livros.append(Livro(**d))
                self.atualizar_lista()

    def exportar_pdf(self):
        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")])
        if path:
            from reportlab.lib.units import inch
            c = canvas.Canvas(path, pagesize=letter)
            width, height = letter
            margin = 50
            y = height - margin

            for livro in self.livros:
                if y < 150:
                    c.showPage()
                    y = height - margin

                # Imagem
                img_x = margin
                img_y = y - 120  # altura da imagem

                if livro.imagem_path and os.path.exists(livro.imagem_path):
                    try:
                        imagem = Image.open(livro.imagem_path)
                        imagem.thumbnail((80, 120), Image.LANCZOS)
                        img_temp_path = "temp_pdf_img.jpg"
                        imagem.save(img_temp_path)
                        c.drawImage(img_temp_path, img_x, img_y, width=80, height=120)
                    except Exception as e:
                        print(f"Erro ao adicionar imagem no PDF: {e}")

                # Texto ao lado da imagem
                text_x = img_x + 100
                text_y = y - 20

                c.setFont("Helvetica-Bold", 14)
                c.drawString(text_x, text_y, livro.titulo)

                c.setFont("Helvetica", 12)
                c.drawString(text_x, text_y - 20, f"Autor: {livro.autor}")

                c.setFont("Helvetica", 11)
                status_texto = livro.status_leitura if hasattr(livro, 'status_leitura') else "Lido" if getattr(livro, 'lido', False) else "Não lido"
                status_text = f"Status: {status_texto} | Nota: {livro.nota} | Tipo: {livro.tipo}"
                c.drawString(text_x, text_y - 40, status_text)

                y -= 140  # Espaço entre livros
    
            c.save()

    def exportar_excel(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
        if path:
            wb = Workbook()
            sheet = wb.active
            sheet.append(["Título", "Autor", "Status", "Nota", "Tipo", "Imagem"])
            for livro in self.livros:
                status = livro.status_leitura if hasattr(livro, 'status_leitura') else "Lido" if getattr(livro, 'lido', False) else "Não lido"
                sheet.append([livro.titulo, livro.autor, status, livro.nota, livro.tipo, livro.imagem_path])
            wb.save(path)

    def exportar_word(self):
        path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Files", "*.docx")])
        if path:
            doc = Document()
            for livro in self.livros:
                doc.add_paragraph(f"Título: {livro.titulo}")
                doc.add_paragraph(f"Autor: {livro.autor}")
                status = livro.status_leitura if hasattr(livro, 'status_leitura') else "Lido" if getattr(livro, 'lido', False) else "Não lido"
                doc.add_paragraph(f"Status: {status}")
                doc.add_paragraph(f"Nota: {livro.nota}")
                doc.add_paragraph(f"Tipo: {livro.tipo}")
                if livro.imagem_path:
                    doc.add_paragraph(f"Imagem: {livro.imagem_path}")
                doc.add_paragraph("")
            doc.save(path)


if __name__ == "__main__":
    app = BibliotecaApp()
    app.mainloop()
